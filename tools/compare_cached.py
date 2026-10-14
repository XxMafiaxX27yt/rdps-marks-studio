"""Compare the values Python cached in a workbook with the values Excel computed.
Usage: python compare_cached.py ours.xlsx excel_recalculated.xlsx"""
import sys, math
import openpyxl

a = openpyxl.load_workbook(sys.argv[1], data_only=True)
b = openpyxl.load_workbook(sys.argv[2], data_only=True)
f = openpyxl.load_workbook(sys.argv[1], data_only=False)
bad = 0
checked = 0
for ws in f.worksheets:
    wa, wb_ = a[ws.title], b[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and c.value.startswith("=")):
                continue
            va, vb = wa[c.coordinate].value, wb_[c.coordinate].value
            checked += 1
            same = va == vb
            if not same and isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                same = math.isclose(va, vb, rel_tol=1e-9, abs_tol=1e-6)
            if not same and (va in (None, "") and vb in (None, "")):
                same = True
            if not same and isinstance(va, str) and isinstance(vb, str) and va.startswith("#") and vb.startswith("#"):
                same = True
            if not same:
                bad += 1
                if bad <= 40:
                    print(f"{ws.title}!{c.coordinate}: python={va!r} excel={vb!r}  formula={c.value[:120]}")
print(f"checked {checked} formulas, {bad} differ")
