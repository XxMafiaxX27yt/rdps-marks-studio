"""Version 2: Classes I-XII, internal marks, the Language column, teachers, and the wider Ask."""
import json
import os
import sys
from decimal import Decimal

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "core"))

from markstudio.config import Config
from markstudio import build, calc, verify, api
from markstudio.layout import Vocab


def run_csv(*files):
    cfg = Config()
    p = build.Project(cfg)
    for name, text in files:
        p.add_file(name, text.encode())
    ds = p.build()
    res = calc.compute(ds)
    verify.check(ds, res)
    return ds, res


def row(res, ds, name, exam, subject):
    for r in res.rows:
        if ds.students[r.sid].name == name and r.exam == exam and r.subject == subject:
            return r
    return None


def test_class_names_every_way():
    v = Vocab(Config())
    for text, want in [("VI-A", ("VI", "A")), ("6th B", ("VI", "B")), ("Class 9 C", ("IX", "C")), ("X", ("X", "")),
                       ("11a", ("XI", "A")), ("xii-b", ("XII", "B")), ("Class 3", ("III", "")), ("10", ("X", ""))]:
        assert v.class_section(text) == want, text
    assert v.find_class("Sheet 1") is None and v.find_class("Term I") is None


def test_marks_only_counts_out_of_written_paper():
    ds, res = run_csv(("Class VI Half Yearly.csv", "Roll,Name,Section,English,Maths\n1,Riya Sen,A,62,40\n"))
    r = row(res, ds, "Riya Sen", "HY", "301")
    assert r.status == "OK" and r.written_only and r.max == Decimal(80) and r.pct == Decimal("77.5")
    assert not any(i.severity == "error" for i in ds.issues)


def test_internal_parts_and_pt_part_from_pt1():
    pt1 = "Roll,Name,Section,English\n1,Riya Sen,A,17\n"
    hy = "Roll,Name,Section,English Written,English Portfolio,English Project,English Notebook\n1,Riya Sen,A,62,4,5,3\n"
    ds, res = run_csv(("Class VI PT1.csv", pt1), ("Class VI Half Yearly.csv", hy))
    r = row(res, ds, "Riya Sen", "HY", "301")
    # PT part = 17/20 x 5 = 4.25; internal = 4.25 + 4 + 5 + 3 = 16.25
    assert r.pr == Decimal("16.25") and r.pr_max == 20 and r.obtained == Decimal("78.25") and r.pct == Decimal("78.25")


def test_one_internal_total_column():
    hy = "Roll,Name,Section,Maths,Maths Internal\n1,Riya Sen,A,50,15\n"
    ds, res = run_csv(("Class IX Half Yearly.csv", hy))
    r = row(res, ds, "Riya Sen", "HY", "041")
    assert r.th == 50 and r.pr == 15 and r.max == 100 and r.pct == Decimal(65)


def test_pt_maximum_depends_on_class():
    ds, res = run_csv(("Class VIII PT1.csv", "Roll,Name,Section,Maths\n1,Om Rao,A,25\n"),
                      ("Class IX PT1.csv", "Roll,Name,Section,Maths\n1,Ira Das,A,25\n"))
    assert row(res, ds, "Om Rao", "PT1", "041").status == "INVALID"           # 25 > 20 in Class VIII
    assert row(res, ds, "Ira Das", "PT1", "041").pct == Decimal("62.5")       # 25 / 40 in Class IX


def test_language_column_uses_each_students_language():
    students = "Roll,Name,Section,Language\n1,Riya Sen,A,French\n2,Om Rao,A,Sanskrit\n"
    hy = "Roll,Name,Section,English,Language\n1,Riya Sen,A,60,64\n2,Om Rao,A,55,71\n"
    ds, res = run_csv(("Class VI Students.csv", students), ("Class VI Half Yearly.csv", hy))
    assert row(res, ds, "Riya Sen", "HY", "018").th == 64
    assert row(res, ds, "Om Rao", "HY", "122").th == 71


def test_same_marks_are_never_duplicates():
    hy = "Roll,Name,Section,English,Maths\n4,Aarav Jain,A,72,65\n9,Diya Roy,A,72,65\n"
    ds, res = run_csv(("Class VII Half Yearly.csv", hy))
    assert len(ds.students) == 2
    assert not any(i.code in ("duplicate_id", "conflict") for i in ds.issues)


def test_ict_rules():
    six = "Roll,Name,Section,English,ICT\n1,Riya Sen,A,60,4\n"
    seven = "Roll,Name,Section,English,ICT\n1,Om Rao,A,60,24\n"
    ds, res = run_csv(("Class VI Half Yearly.csv", six), ("Class VII Half Yearly.csv", seven))
    assert row(res, ds, "Riya Sen", "HY", "402") is None                          # not counted in I-VI
    assert row(res, ds, "Om Rao", "HY", "402").pct == Decimal(80)                 # 24 / 30 from VII


def test_weak_student_needs_support_only_where_below_pass():
    hy = "Roll,Name,Section,English,Maths,Science\n1,A One,A,70,70,70\n2,B Two,A,70,72,74\n3,C Three,A,20,40,42\n"  # English 20/80 = 25%
    ds, res = run_csv(("Class VIII Half Yearly.csv", hy))
    labels = {r.subject: r.label for r in res.rows if ds.students[r.sid].name == "C Three"}
    assert labels["301"] == "Needs attention" and labels["041"] != "Needs attention" and labels["086"] != "Needs attention"


def test_best_subject_is_not_a_focus_area():
    hy = "Roll,Name,Section,English,Maths,Science\n1,A One,A,78,79,77\n2,B Two,A,79,78,80\n3,C Three,A,60,48,47\n"
    ds, res = run_csv(("Class VIII Half Yearly.csv", hy))
    sid = next(s.sid for s in ds.students.values() if s.name == "C Three")
    rem = calc.remarks(res, sid, "HY")
    best = rem[0]
    assert best.startswith("Best subject: English")
    assert not any(x.startswith("Focus area: English") for x in rem)


def test_teacher_analysis_and_fair_check():
    from markstudio.teachers import analyse, normalise
    hy = ("Roll,Name,Section,English,Maths\n1,A One,A,80,60\n2,B Two,A,70,50\n"
          "1,C Three,B,60,62\n2,D Four,B,50,52\n")
    ds, res = run_csv(("Class X Half Yearly.csv", hy))
    cfg = ds.cfg
    ts = normalise([{"name": "Ms. Maths A", "subject": "Maths", "classes": "X-A"},
                    {"name": "Mr. Maths B", "subject": "041", "classes": "X-B"}], cfg)
    an = analyse(res, ts)
    a, b = an["cards"]
    assert a["classes"] == ["X-A"] and b["classes"] == ["X-B"]
    # marks out of 80: section A Maths is 12.5 below the students' own averages; section B 1.25 above
    assert a["overall"]["vs_own"] == Decimal("-12.5") and b["overall"]["vs_own"] == Decimal("1.25")
    # side by side, higher average first (B: 71.25% vs A: 68.75%) - even though A's students are stronger overall
    assert [r["name"] for r in an["comparison"][0]["rows"]] == ["Mr. Maths B", "Ms. Maths A"]


def test_ask_understands_every_class_and_new_questions():
    api.init()
    api.add_sample(False)
    api.analyze()
    cases = {
        "result of class 10": ("summary", "X"),
        "who topped class 6": ("top", "VI"),
        "section toppers class 12": ("toppers_by", "XII"),
        "weakest subject of class 9": ("hardest", "IX"),
        "who failed in class 10": ("list", "X"),
        "class 9 mein top 5 bachhe": ("top", "IX"),
        "students with distinction class 12": ("list", "XII"),
        "how many got A1 in class 11": ("count", "XI"),
        "which section is best in class 12": ("average", "XII"),
        "teacher wise result": ("teacher", None),
    }
    for q, (op, cls) in cases.items():
        out = json.loads(api.ask(q))
        assert out["ok"], q
        d = out["data"]
        assert "query" in d, (q, d)
        assert d["query"]["op"] == op, (q, d["query"])
        if cls:
            assert d["query"]["scope"].get("cls") == cls, (q, d["query"])


def test_ask_internal_marks_and_meanings():
    api.init()
    api.add_sample(False)
    api.analyze()
    d = json.loads(api.ask("internal marks of class 6"))["data"]
    assert d["query"]["op"] == "internals" and d["table"]["rows"]
    d = json.loads(api.ask("who has low internal marks in class 9"))["data"]
    assert d["query"].get("level") == "low"
    d = json.loads(api.ask("who all are weak in chemistry"))["data"]
    assert d["query"]["conds"] == [{"level": "low", "subject": "043"}]
    pcts = [float(r[3].split("Chemistry ")[1].split("%")[0]) for r in d["table"]["rows"]]
    assert pcts == sorted(pcts) and all(p < 50 for p in pcts)          # lowest first, all below 50%
    d = json.loads(api.ask("which teachers are not working good"))["data"]
    assert d["query"]["op"] == "teacher" and d["query"].get("order") == "low"


def test_reopened_workbook_then_settings_or_new_files():
    api.init()
    api.add_sample(False)
    api.analyze()
    wb = api.build_workbook_bytes(False)
    assert json.loads(api.load_workbook("saved.xlsx", wb))["ok"]
    st = json.loads(api.get_settings())["data"]
    st["rules"]["pass_percent"] = 40
    assert json.loads(api.set_settings(json.dumps(st)))["ok"]
    a = json.loads(api.analyze())
    assert a["ok"] and a["data"]["students"] == 364
    api.add_file("Class VII Half Yearly.csv", b"Roll,Name,Section,English,Maths\n1,Riya Sen,A,62,40\n")
    a = json.loads(api.analyze())["data"]
    assert a["students"] == 1 and not a["restored"]


def test_remove_teachers_completely():
    api.init()
    api.add_sample(False)
    api.analyze()
    lst = json.loads(api.teachers_list())["data"]["from_files"]
    t = lst[0]
    json.loads(api.remove_file_teacher(t["name"], t["subject"]))
    names = [c["name"] for c in json.loads(api.teacher_analysis())["data"]["cards"]]
    assert t["name"] not in names
    json.loads(api.remove_all_teachers())
    assert json.loads(api.teacher_analysis())["data"]["cards"] == []


def test_board_teachers_on_sample_gazette():
    import io
    import openpyxl
    api.init()
    assert json.loads(api.add_sample_gazette())["ok"]
    api.set_teachers(json.dumps([{"name": "Ms. Test One", "subject": "Physics", "classes": "XII"},
                                 {"name": "Mr. Test Two", "subject": "English", "classes": "XII-A, XII-B"},
                                 {"name": "Ms. Test Three", "subject": "English", "classes": "XII-C"}]))
    t = json.loads(api.board_summary())["data"]["teachers"]
    names = {c["subject_name"]: c for c in t["cards"]}
    # no school marks with sections, so the two English teachers are shown together
    assert names["English"]["names"] == ["Mr. Test Two", "Ms. Test Three"] and names["English"]["note"]
    assert names["Physics"]["fair"]["verdict"] == "Only teacher of this subject"
    assert "Mathematics" in [u["name"] for u in t["untaught"]]
    d = json.loads(api.ask("which teachers are not doing well in boards"))["data"]
    assert d["query"] == {"op": "board_teachers", "order": "low"} and d["table"]["rows"]
    wb = openpyxl.load_workbook(io.BytesIO(api.board_workbook_bytes()))
    assert "Teachers" in wb.sheetnames


def test_board_teachers_on_sample_gazette():
    import io
    import openpyxl
    api.init()
    assert json.loads(api.add_sample_gazette())["ok"]
    api.set_teachers(json.dumps([{"name": "Ms. Test One", "subject": "Physics", "classes": "XII"},
                                 {"name": "Mr. Test Two", "subject": "English", "classes": "XII-A, XII-B"},
                                 {"name": "Ms. Test Three", "subject": "English", "classes": "XII-C"}]))
    t = json.loads(api.board_summary())["data"]["teachers"]
    names = {c["subject_name"]: c for c in t["cards"]}
    # no school marks with sections, so the two English teachers are shown together
    assert names["English"]["names"] == ["Mr. Test Two", "Ms. Test Three"] and names["English"]["note"]
    assert names["Physics"]["fair"]["verdict"] == "Only teacher of this subject"
    assert "Mathematics" in [u["name"] for u in t["untaught"]]
    d = json.loads(api.ask("which teachers are not doing well in boards"))["data"]
    assert d["query"] == {"op": "board_teachers", "order": "low"} and d["table"]["rows"]
    wb = openpyxl.load_workbook(io.BytesIO(api.board_workbook_bytes()))
    assert "Teachers" in wb.sheetnames


# ---------------------------------------------------------------- board tables, class lists, each teacher on their own
TABLE_TXT = """Class XII Board Examination 2025 - Board Results (made-up)
Each subject out of 100.

Roll Admn No  Student Name              Eng  Phy  Chem  Math  CS  Total   Perc%  Grd  Result
--------------------------------------------------------------------------------------------
1    9901     Riya Sen                   89   91    85    72  93    430    86.0   A2  PASS
2    C-17     Om Prakash Rao             64   68    80    81  63    356    71.2   B1  PASS
3    9905     Asha Devi                  30   45    50    41  52    999    43.6   C2  COMP
"""


def test_board_results_in_the_schools_own_table():
    api.init()
    info = json.loads(api.add_board_file("XII Board Results.txt", TABLE_TXT.encode()))["data"]
    assert info["table"] and info["students"] == 3 and info["cls"] == "XII"
    g = api.S.project.gazettes[-1]
    om = next(x for x in g.students if x.name == "Om Prakash Rao")
    assert om.adm_no == "C-17" and [x["marks"] for x in om.subjects] == [64, 68, 80, 81, 63]
    asha = next(x for x in g.students if x.name == "Asha Devi")
    assert asha.result == "COMP" and asha.comp == ["301"]
    codes = {i.code for i in g.issues}
    assert "file_total_mismatch" in codes and "board_no_grades" in codes        # 999 is not the sum
    b = json.loads(api.board_summary())["data"]["boards"][0]
    eng = next(x for x in b["subjects"] if x["code"] == "301")
    assert eng["failed"] == 1 and eng["bands"][0] == ["90-100", 0]


def test_a_school_file_with_a_board_roll_column_stays_a_school_file():
    from markstudio import boardtable, readers
    from markstudio.sample import sample_files
    for name, data in sample_files(Config()):
        assert not boardtable.looks_like_board(name, readers.read_file(name, data)), name
    assert boardtable.looks_like_board("XII_Board_Marks.txt", readers.read_file("x.txt", TABLE_TXT.encode()))
    assert not boardtable.looks_like_board("Pre Board XII.csv", readers.read_file("x.csv", b"Name,English\nA,1\n"))


def test_class_list_from_pdf_words_with_wrapped_name_and_footer():
    from markstudio import readers, classlist
    rows = [["Class & Sec: XII B (2025-26)"], [], ["SNo.", "Roll No.", "Admn. No.", "Student's Name", "Remarks"],
            ["1", "1", "9901", "Riya Sen"], ["2", "2", "9902", "Om Prakash"], ["", "", "", "Rao"],
            ["3", "3", "C-17", "Asha Devi"], ["10:21 AM Page 1 of 1"]]
    xs = [40, 110, 170, 260, 480]
    words = [{"text": t, "x0": xs[c], "x1": xs[c] + 6 * len(t), "y0": 20 * r, "y1": 20 * r + 10}
             for r, row in enumerate(rows) for c, t in enumerate(row) if t]
    sheets, _ = readers.read_pages("XII B list.pdf", [{"page": 1, "words": words, "ocr": False}])
    cl = classlist.read("XII B list.pdf", sheets, Config(), Vocab(Config()))
    assert len(cl) == 1 and (cl[0]["cls"], cl[0]["section"]) == ("XII", "B")
    assert [s["name"] for s in cl[0]["students"]] == ["Riya Sen", "Om Prakash Rao", "Asha Devi"]
    assert cl[0]["students"][2]["adm_no"] == "C-17"


def test_each_teacher_on_their_own_sections():
    import io
    import openpyxl
    api.init()
    out = json.loads(api.add_sample_board_demo())["data"]
    assert [(c["cls"], c["section"], c["count"]) for c in out["class_lists"]] == [("XII", "A", 30), ("XII", "B", 30), ("XII", "C", 30)]
    bt = json.loads(api.board_teachers())["data"]
    p = bt["placement"]
    assert bt["split"] and p["placed"] == 90 and p["by"]["admission no"] == 60 and len(p["check"]) == 1
    by = {c["name"]: c for c in bt["cards"]}
    # the made-up data has Physics XII-A 5 marks better and XII-B / XII-C 2 marks lower
    assert by["Ms. Neha Sinha"]["candidates"] == 30 and by["Ms. Neha Sinha"]["fair"]["verdict"] == "Clear positive impact"
    assert by["Mr. Rohan Das"]["candidates"] == 60 and by["Mr. Rohan Das"]["fair"]["verdict"] == "Worth a conversation"
    assert by["Ms. Priya Menon"]["fair"]["verdict"] == "Only teacher of this subject"
    v = by["Ms. Neha Sinha"]["viz"]
    assert sum(v["split"].values()) == 30 and len(v["points"]) == 30 and v["bands_are"] == "marks"
    wb = openpyxl.load_workbook(io.BytesIO(api.board_workbook_bytes()))
    assert "Teachers" in wb.sheetnames and "Students & Sections" in wb.sheetnames
    assert len(wb["Teachers"]._charts) == 2
    d = json.loads(api.ask("which teachers are not doing well in boards"))["data"]
    assert d["query"]["op"] == "board_teachers" and d["table"]["rows"][0][1] == "Mr. Rohan Das"


def test_class_lists_give_school_students_their_section():
    import io
    import xlsxwriter
    api.init()
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf)
    ws = wb.add_worksheet("XI C Roll List")
    for r, row in enumerate([["Class & Sec: XI C"], [], ["Roll No.", "Admn. No.", "Student's Name"], [1, 501, "Riya Sen"], [2, 502, "Om Rao"]]):
        ws.write_row(r, 0, row)
    wb.close()
    api.add_class_list("XI C list.xlsx", buf.getvalue())
    api.add_file("Class XI Half Yearly.csv", b"Name,English,Maths\nRiya Sen,60,55\nOm Rao,70,65\n")
    a = json.loads(api.analyze())["data"]
    assert {s.section for s in api.S.ds.students.values()} == {"C"}


def test_total_percentage_grade_result_columns_are_understood():
    api.init()
    api.add_file("Class XII-B Half Yearly.csv", b"Roll,Name,Eng,Maths,Total,Perc%,Grd,Result\n1,Asha Rao,70,65,135,67.5,B1,PASS\n")
    st = json.loads(api.state())["data"]
    roles = {c["header"]: c["role"] for c in st["files"][0]["sheets"][0]["columns"]}
    assert roles["Total"] == "file_total" and roles["Perc%"] == "file_percent"
    assert roles["Grd"] == "file_grade" and roles["Result"] == "file_result"


def test_any_file_is_sorted_by_itself():
    from markstudio.sample import sample_files, sample_board_table, sample_class_lists, sample_gazette
    api.init()
    cfg = api.S.cfg
    kinds = [json.loads(api.add_any(n, d))["data"]["kind"] for n, d in sample_files(cfg)]
    assert kinds == ["school"] * 5
    n, t = sample_board_table(cfg)
    i = json.loads(api.add_any(n, t.encode()))["data"]
    assert (i["kind"], i["cls"], i["what"]) == ("board", "XII", "table")
    i = json.loads(api.add_any("gazette.txt", sample_gazette(cfg).encode()))["data"]
    assert (i["kind"], i["cls"], i["what"]) == ("board", "XII", "cbse")
    n, d = sample_class_lists(cfg)[0]
    i = json.loads(api.add_any(n, d))["data"]
    assert i["kind"] == "class_list" and i["lists"][0]["section"] == "A"
    # a plain marks table with no exam and no 'board' in it: the tool asks (or follows the page it was dropped on)
    plain = b"Roll,Name,English,Physics,Chemistry,Maths,CS\n1,Asha Rao,80,70,75,66,90\n"
    assert json.loads(api.add_any("Class 12 A.csv", plain))["data"]["kind"] == "ask"
    assert json.loads(api.add_any("Class 12 A.csv", plain, "", "board"))["data"]["kind"] == "board"
    i = json.loads(api.add_any("Class 10 B.csv", b"Roll,Name,English,Maths,Science\n1,Om Das,80,70,75\n", "board_X"))["data"]
    assert (i["kind"], i["cls"]) == ("board", "X")
    assert json.loads(api.add_any("Class 6 Half Yearly.csv", b"Roll,Name,English,Maths\n1,Om Das,60,70\n"))["data"]["kind"] == "school"
    st = json.loads(api.state())["data"]
    assert len(st["gazettes"]) == 4 and len(st["class_lists"]) == 1


# ---------------------------------------------------------------- Board page v4: streams, highlights, toppers, one workbook per class
CLASS_X_CSV = (b"Roll,Name,English,Hindi,Maths,Science,Social Science\n"
               b"1,Om Das,80,70,75,66,90\n2,Asha Rao,91,88,95,93,90\n3,Kabir Sen,45,52,30,41,58\n")


def test_board_page_numbers_on_the_made_up_gazette():
    from markstudio.board import STREAM_ORDER
    api.init()
    assert json.loads(api.add_sample_gazette())["ok"]
    b = json.loads(api.board_summary())["data"]["boards"][0]
    assert b["label"] == "Class XII · 2025" and b["key"] == "0"
    # streams: every student in exactly one, toppers sorted
    assert [s["stream"] for s in b["streams"]] == [x for x in STREAM_ORDER if x in {s["stream"] for s in b["streams"]}]
    assert sum(s["n"] for s in b["streams"]) == b["candidates"]
    for s in b["streams"]:
        assert [t["best5"] for t in s["top"]] == sorted((t["best5"] for t in s["top"]), reverse=True)
    # subjects: middle mark and middle half, how the same students did elsewhere
    for s in b["subjects"]:
        assert s["q1"] <= s["median"] <= s["q3"] and s["min"] <= s["q1"] and s["q3"] <= s["max"]
    eng = next(s for s in b["subjects"] if s["code"] == "301")
    assert eng["vs_own"] is not None and eng["vs_own_n"] == b["candidates"]
    # every student: stream, the 5 subjects counted, rank
    for st in b["students"]:
        assert st["stream"] and len(st["best5_subjects"]) == 5 and "301" in st["best5_subjects"]
    assert b["highlights"][0].startswith(f"{b['passed']} of {b['candidates']} students passed")
    assert any(h.startswith("Top scorer:") for h in b["highlights"])


def test_quartiles_are_excels_and_the_other_subjects_comparison():
    from markstudio.board import _quartile, subject_vs_own

    class St:
        def __init__(self, marks):
            self.subjects = [{"code": c, "marks": m} for c, m in marks]

    assert [float(_quartile([1, 2, 3, 4], k)) for k in (1, 2, 3)] == [1.75, 2.5, 3.25]   # = QUARTILE.INC / MEDIAN
    assert _quartile([7], 2) == 7 and _quartile([], 2) is None
    rows = [{"s": St([("A", 90), ("B", 70), ("C", 80)])}, {"s": St([("A", 60), ("B", 50), ("C", 70)])}]
    d = subject_vs_own(rows)
    assert float(d["A"][0]) == 7.5 and d["A"][1] == 2        # (90-75 + 60-60) / 2
    assert float(d["B"][0]) == -15.0


def test_the_same_board_file_twice_is_kept_once():
    from markstudio.sample import sample_gazette
    api.init()
    t = sample_gazette(api.S.cfg).encode()
    api.add_any("gazette.txt", t)
    api.add_any("gazette copy.txt", t)
    assert len(api.S.project.gazettes) == 1
    # two different lists of the same class (no roll numbers) are both kept
    api.add_any("XII A board.csv", b"Name,English,Physics,Chemistry,Maths,PE\nAsha Rao,80,70,75,66,90\n", "board_XII")
    api.add_any("XII B board.csv", b"Name,English,Physics,Chemistry,Maths,PE\nOm Das,80,70,75,66,90\n", "board_XII")
    assert len(api.S.project.gazettes) == 3


def test_class_x_and_xii_each_get_their_own_view_and_workbook():
    import io
    import openpyxl
    api.init()
    api.add_sample_gazette()
    i = json.loads(api.add_any("Class 10 Board Results.csv", CLASS_X_CSV, "board_X"))["data"]
    assert (i["kind"], i["cls"]) == ("board", "X")
    boards = json.loads(api.board_summary())["data"]["boards"]
    assert [b["label"].split(" · ")[0] for b in boards] == ["Class XII", "Class X"]
    x = boards[1]
    assert x["streams"] == [] and all(s["stream"] == "" for s in x["students"])
    assert [s["name"] for s in sorted(x["students"], key=lambda s: s["rank"])][0] == "Asha Rao"
    for cls, key in (("XII", ""), ("X", ""), ("", "1")):
        wb = openpyxl.load_workbook(io.BytesIO(api.board_workbook_bytes(cls, key)))
        want = "X" if (cls == "X" or key == "1") else "XII"
        assert wb["Board Dashboard"]["C1"].value.endswith(f"Class {want}")
        assert ("Streams" in wb.sheetnames) == (want == "XII")
        assert api.board_workbook_name(cls, key).endswith(f"Class {want}.xlsx")


def test_board_workbook_results_in_merit_order_with_formulas():
    import io
    import openpyxl
    api.init()
    api.add_sample_gazette()
    wb = openpyxl.load_workbook(io.BytesIO(api.board_workbook_bytes()))
    assert wb.sheetnames[:5] == ["Board Dashboard", "Results", "Subjects", "Toppers", "Streams"]
    ws = wb["Results"]
    heads = [c.value for c in ws[5]]
    b5, rk = heads.index("Best 5 %") + 1, 1
    first = ws.cell(6, b5).value
    assert first.startswith("=SUM(") and first.endswith(")/5")                  # the 5 counted marks
    assert ws.cell(6, rk).value.startswith('=IF(N(') and "RANK(" in ws.cell(6, rk).value
    cached = openpyxl.load_workbook(io.BytesIO(api.board_workbook_bytes()), data_only=True)["Results"]
    vals = [cached.cell(r, b5).value for r in range(6, 6 + 60)]
    assert vals == sorted(vals, reverse=True)                                    # merit order
    s = wb["Subjects"]
    assert [c.value for c in s[5]][:13][-1] == "vs own other subjects"
    assert s["D6"].value.startswith("=IFERROR(AVERAGE('Results'!") and s["E6"].value.startswith("=IFERROR(MEDIAN(")
    d = wb["Board Dashboard"]
    assert d["B7"].value.startswith("=COUNTA('Results'!") and str(d["Q7"].value).startswith('=COUNTIF(')
    t = [c.value for row in wb["Toppers"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "Top three" in t and "Stream toppers" in t and "Subject toppers" in t


def test_board_analysis_template_round_trip_and_the_made_up_example():
    import io
    import openpyxl
    api.init()
    wb = openpyxl.load_workbook(io.BytesIO(api.board_template_bytes()))
    assert wb.sheetnames == ["How to fill", "Students", "Teachers"] and wb["Students"].max_row == 1
    ws, wt = wb["Students"], wb["Teachers"]
    ws.append(["99100001", "4500", "Asha Rao", "XII", "A"])
    ws.append(["99100002", "4501", "Om Das", "X", "C"])
    wt.append(["Ms. Example One", "Physics", "XII-A, XII-B"])
    buf = io.BytesIO()
    wb.save(buf)
    i = json.loads(api.add_any("Board analysis.xlsx", buf.getvalue(), "", "board"))["data"]
    assert i["kind"] == "class_list"                                   # the Teachers sheet is not read as students
    studs = api.S.project.class_lists[0]["students"]
    assert [(s["name"], s["cls"], s["section"], s["board_roll"]) for s in studs] == \
        [("Asha Rao", "XII", "A", "99100001"), ("Om Das", "X", "C", "99100002")]
    assert [(t["name"], t["subject"]) for t in api.S.project.file_teachers] == [("Ms. Example One", "Physics")]
    # the whole made-up example: every student placed, every subject has a teacher, twice = once
    api.init()
    api.add_sample_board_full()
    api.add_sample_board_full()
    assert len(api.S.project.class_lists) == 1 and len(api.S.project.gazettes) == 1
    t = json.loads(api.board_summary())["data"]["teachers"]
    assert t["placement"]["placed"] == t["placement"]["total"] == 60 and t["split"] and not t["untaught"]
    eng = [c for c in t["cards"] if c["subject_name"] == "English"]
    assert len(eng) == 2 and eng[0]["fair"]["verdict"] != "Only teacher of this subject"
    # the made-up teacher effects show every kind of result
    v = {c["name"]: c["fair"]["verdict"] for c in t["cards"]}
    assert v["Mr. Arjun Malhotra"] == v["Mr. Rohit Bansal"] == "Clear positive impact"
    assert v["Ms. Kiran Mehta"] == "Worth a conversation" and "No clear difference" in v.values()
    ex = openpyxl.load_workbook(io.BytesIO(api.board_template_bytes(True)))
    assert ex["Students"].max_row == 61 and ex["Teachers"].max_row > 10


def test_ask_board_streams_all_rounders_and_hardest():
    api.init()
    api.add_sample_gazette()
    ask = lambda q: json.loads(api.ask(q))["data"]
    assert ask("stream toppers")["query"] == {"op": "board_streams"}
    d = ask("commerce top 3")
    assert d["query"] == {"op": "board_streams", "stream": "Commerce", "n": 3} and len(d["table"]["rows"]) >= 3
    assert ask("who got 90 in every subject")["query"] == {"op": "board_allround", "value": 90}
    assert ask("sab subject me 95 se upar kaun hai")["query"]["value"] == 95
    assert ask("hardest subject in boards")["query"] == {"op": "board_hardest", "order": "hard"}
    assert ask("easiest subject")["query"]["order"] == "easy"
    assert ask("top 5 in physics")["query"] == {"op": "board_top", "n": 5, "subject": "042"}   # no school marks: all board
    # Class X and XII both loaded: the class in the question picks the results
    api.add_any("Class 10 Board Results.csv", CLASS_X_CSV, "board_X")
    d = ask("class 10 board toppers")
    assert d["query"]["cls"] == "X" and d["table"]["rows"][0][1] == "Asha Rao"


def test_class_picked_before_uploading():
    """The School page's class drop-down: a file that does not say its class is that class; one that looks like
    another class (title, class column, subjects the class does not have) gets a question until answered."""
    cfg = Config()
    assert "042" not in cfg.class_subject_codes("X") and "086" in cfg.class_subject_codes("X")
    assert "042" in cfg.class_subject_codes("XII") and "086" not in cfg.class_subject_codes("XII")
    from markstudio.layout import class_parts
    assert "portfolio" in class_parts(cfg, "X") and "portfolio" not in class_parts(cfg, "XII")

    def add(name, text, cls):
        p = build.Project(cfg)
        info = p.add_any(name, text.encode(), cls=cls)
        lay = next(iter(p.layouts.values()))
        return p, info, lay
    plain = "Roll,Name,English,Maths,Science,SSt\n1,Asha,60,70,55,66\n2,Ravi,50,40,45,61\n"
    p, info, lay = add("half yearly marks.csv", plain, "X")
    assert lay["cls"] == "X" and lay["cls_from"] == "picked" and not lay.get("clash") and not p.pending()
    ds = p.build()
    assert {s.cls for s in ds.students.values()} == {"X"}
    p, info, lay = add("half yearly marks.csv", plain, "")
    assert lay["cls"] is None and [x["what"] for x in p.pending()] == ["class"]

    phy = "Roll,Name,English,Physics,Chemistry,Maths\n1,Asha,60,70,55,66\n2,Ravi,50,40,45,61\n"
    p, info, lay = add("half yearly marks.csv", phy, "X")
    assert lay["clash"]["source"] == "subjects" and lay["clash"]["fits"] == ["XI", "XII"]
    assert "Physics and Chemistry are not Class X subjects" in lay["clash"]["why"]
    assert info["clashes"] and "class_clash" in [x["what"] for x in p.pending()]
    p.set_sheet(p.sources[0].id, lay["sheet"], cls="XII")
    lay = p.layouts[(p.sources[0].id, lay["sheet"])]
    assert not lay.get("clash") and lay["cls"] == "XII" and "class_clash" not in [x["what"] for x in p.pending()]
    # worked out again for Class XII: one Physics column could be the written paper (70) or the total (100)
    phys = next(c for c in lay["columns"] if c["subject"] == "042")
    assert not phys["sure"] and "Physics" in phys["why"]

    titled = "Class IX Half Yearly\nRoll,Name,English,Maths\n1,Asha,60,70\n"
    p, info, lay = add("marks.csv", titled, "X")
    assert lay["clash"]["found"] == "IX" and lay["clash"]["source"] == "title"
    p, info, lay = add("marks.csv", titled, "IX")
    assert not lay.get("clash") and lay["cls"] == "IX"

    col = "Roll,Name,Class,English,Maths\n1,Asha,IX-A,60,70\n2,Ravi,IX-B,50,40\n"
    p, info, lay = add("half yearly marks.csv", col, "X")
    assert lay["clash"]["source"] == "column" and lay["clash"]["seen"] == ["IX"]
    p.set_sheet(p.sources[0].id, lay["sheet"], skip=True)
    assert not p.pending()

    # the class picked for this upload beats a layout remembered from another class's file
    p, info, lay = add("half yearly marks.csv", plain, "VII")
    p.confirm_sheet(p.sources[0].id, lay["sheet"])
    p.add_any("half yearly marks b.csv", plain.encode(), cls="VIII")
    second = p.layouts[(p.sources[1].id, p.sources[1].sheets[0].name)]
    assert second["remembered"] and second["cls"] == "VIII"

    api.init()
    info = json.loads(api.class_info("X"))["data"]
    assert info and "Physics" not in [s["name"] for s in info["subjects"]] and "Hindi" in info["languages"]


def test_type_marks_screen():
    """Marks typed on the website: names from the class list, only the parts ticked, empty students left out,
    typing the same class / exam / subject again replaces the first copy, and results are made from it."""
    api.init()
    api.add_any("Class VI-A list.csv", b"Roll No,Admission No,Name\n1,9001,Asha Test\n2,9002,Ravi Test\n3,9003,Meena Test\n",
                "class_list", "", "")
    setup = json.loads(api.type_setup("VI", "A"))["data"]
    assert [s["name"] for s in setup["students"]] == ["Asha Test", "Ravi Test", "Meena Test"] and setup["sections"] == ["A"]
    hy = {c["part"]: (c["max"], c["on"]) for c in setup["columns"]["HY|301"]}
    assert hy == {"theory": (80, True), "portfolio": (5, False), "project": (5, False), "notebook": (5, False)}
    assert setup["columns"]["HY|402"] == []                      # ICT is not counted in Class VI
    xii = json.loads(api.type_setup("XII", ""))["data"]["columns"]["HY|042"]
    assert [(c["part"], c["max"], c["on"]) for c in xii] == [("theory", 70, True), ("practical", 30, True)]

    def typed(values, subject="301"):
        rows = [{"roll": s["roll"], "adm_no": s["adm_no"], "name": s["name"], "section": "A",
                 "values": {"theory": v} if v is not None else {}} for s, v in zip(setup["students"], values)]
        return json.loads(api.add_typed(json.dumps({"cls": "VI", "section": "A", "exam": "HY", "subject": subject,
                                                    "parts": ["theory"], "rows": rows, "teacher": ""})))
    out = typed(["66", "AB", None])
    assert out["ok"] and out["data"]["students"] == 2
    out = typed(["70", "AB", None])                               # typed again: replaces, not added twice
    st = json.loads(api.state())["data"]
    assert len(st["files"]) == 1 and st["files"][0]["typed"]["count"] == 2 and not st["pending"]
    typed(["50", "60", "40"], subject="041")
    res = json.loads(api.analyze())["data"]
    assert res["students"] == 3 and not res["counts"]["error"]
    m = json.loads(api.matrix("VI", "HY", "All"))["data"]
    asha = next(r for r in m["rows"] if r["name"] == "Asha Test")
    assert [c["pct"] for c in asha["cells"] if c] == [87.5, 62.5]      # 70 / 80 and 50 / 80
    again = json.loads(api.typed_get(st["files"][0]["id"]))["data"]
    assert again["rows"][0]["values"] == {"theory": "70"}
    assert not json.loads(api.add_typed(json.dumps({"cls": "VI", "section": "A", "exam": "HY", "subject": "301",
                                                    "parts": ["theory"], "rows": [], "teacher": ""})))["ok"]


def test_ask_question_bank():
    """Every question in tests/ask_bank.json (written by a local AI model from clear questions with blanks, filled with
    many ways of writing a class, subject, name ...) still gives the same query as the clear question."""
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))
    import ask_bank
    bank = json.load(open(os.path.join(os.path.dirname(__file__), "ask_bank.json"), encoding="utf-8"))
    assert sum(len(v) for v in bank["phrasings"].values()) > 500
    bad = ask_bank.check_bank(bank)
    assert not bad, bad[:5]
