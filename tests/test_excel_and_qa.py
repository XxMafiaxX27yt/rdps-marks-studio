"""Workbook, macros, word commands and Q&A tests (spec sections 11, 14, 16)."""
import io
import json
import os
import socket
import sys
from decimal import Decimal

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "core"))

from markstudio.config import Config
from markstudio import build, calc, verify, sample, api
from markstudio.nlu import Understander
from markstudio.query import Engine
from markstudio.excel import vba
from markstudio.excel.school import write_workbook


@pytest.fixture(scope="module")
def data():
    cfg = Config()
    p = build.Project(cfg)
    for n, d in sample.sample_files(cfg, mistakes=False):
        p.add_file(n, d)
    ds = p.build()
    res = calc.compute(ds)
    verify.check(ds, res)
    return ds, res


@pytest.fixture(autouse=True)
def no_internet(monkeypatch):
    """The whole tool must work with no internet: any network use fails the test."""
    def boom(*a, **k):
        raise AssertionError("network access attempted")
    monkeypatch.setattr(socket, "socket", boom)
    monkeypatch.setattr(socket, "create_connection", boom)


def ask(data, q, force=None):
    ds, res = data
    U = Understander(ds, res)
    parsed = U.parse(q, force=force)
    if "query" not in parsed:
        return parsed, None
    return parsed, Engine(ds, res).run(parsed["query"])


# ---------------------------------------------------------------- workbook
def test_workbook_opens_and_has_sheets(data):
    import openpyxl
    ds, res = data
    b = write_workbook(ds, res)
    wb = openpyxl.load_workbook(io.BytesIO(b))
    for name in ["Dashboard", "Report Card", "Subject Analysis", "Marks Analysis", "Progress", "Overview",
                 "Marks VI", "Marks XI PCM", "Internals VI", "Subject Results", "Student Master", "Validation Report",
                 "Config", "Source Data"]:
        assert name in wb.sheetnames
    assert wb["_meta"].sheet_state == "hidden"


def test_workbook_cached_values_match_python(data):
    import openpyxl
    ds, res = data
    b = write_workbook(ds, res)
    wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
    ws = wb["Subject Results"]
    head = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(head)}
    cfg = ds.cfg
    by_key = {f"{r.sid}|{cfg.exam_name(r.exam)}|{cfg.subject_name(r.subject)}": r for r in res.rows}
    n = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = by_key[r[ci["Key"]]]
        if row.pct is not None:
            assert abs(r[ci["%"]] - float(row.pct)) < 1e-9
        assert r[ci["Label"]] == row.label and r[ci["Reason"]] == row.reason
        n += 1
    assert n == len(res.rows)


def test_uploaded_text_never_becomes_a_formula():
    cfg = Config()
    p = build.Project(cfg)
    p.add_file("Half Yearly XI.csv", b"Roll,Name,Class,Section,Maths\n1,=HYPERLINK(\"x\"),XI,A,60\n")
    ds = p.build()
    res = calc.compute(ds)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(write_workbook(ds, res)))
    ws = wb["Student Master"]
    assert ws["E2"].value == '=HYPERLINK("x")' and ws["E2"].data_type == "s"


def test_macros_project_is_valid(data):
    from oletools.olevba import VBA_Parser
    ds, res = data
    b = write_workbook(ds, res, vba=True)
    vp = VBA_Parser("t.xlsm", data=b)
    names = [n for (_, _, n, _) in vp.extract_macros()]
    assert "RDPSMacros.bas" in names and "ThisWorkbook.cls" in names


def test_vba_compression_roundtrip():
    src = vba.module_source().encode("cp1252")
    for d in [src, src * 3, b"", b"x" * 5000, os.urandom(4096)]:
        assert vba.decompress(vba.compress(d)) == d


def test_reupload_workbook_with_excel_edit(data):
    import openpyxl
    ds, res = data
    api.S = api.Session()
    api.S.ds, api.S.res = ds, res
    b = write_workbook(ds, res)
    wb = openpyxl.load_workbook(io.BytesIO(b))
    ws = wb["Marks XI PCM"]
    # find first numeric mark cell and change it
    target = None
    for row in ws.iter_rows(min_row=5, max_row=6):
        for c in row[6:]:
            if isinstance(c.value, (int, float)):
                target = c
                break
        if target:
            break
    old = target.value
    target.value = 1 if old != 1 else 2
    buf = io.BytesIO()
    wb.save(buf)
    out = json.loads(api.load_workbook("edited.xlsx", buf.getvalue()))
    assert out["ok"], out
    assert out["data"]["changes_in_excel"] == 1
    assert any(i["code"] == "changed_in_excel" for i in out["data"]["issues"])


# ---------------------------------------------------------------- Q&A
def test_top_n_matches_hand_sort(data):
    ds, res = data
    parsed, ans = ask(data, "top 5 in class 11")
    assert parsed["query"]["op"] == "top" and parsed["query"]["n"] == 5
    want = sorted([er.overall for (sid, e), er in res.exam_results.items()
                   if ds.students[sid].cls == "XI" and e == ds.exams_for("XI")[-1] and er.overall is not None], reverse=True)[:5]
    got = [Decimal(str(r[4])) for r in ans["table"]["rows"][:5]]
    assert [round(x, 2) for x in want] == [round(x, 2) for x in got]


def test_weak_and_strong_combination(data):
    ds, res = data
    parsed, ans = ask(data, "Who is weak in Maths but strong in Physics in class XI?")
    conds = parsed["query"]["conds"]
    assert {"level": "low", "subject": "041"} in conds and {"level": "high", "subject": "042"} in conds
    ex = ds.exams_for("XI")[-1]
    want = set()
    for (sid, e), er in res.exam_results.items():
        if ds.students[sid].cls != "XI" or e != ex:
            continue
        by = {r.subject: r for r in er.rows if r.pct is not None}
        m, p = by.get("041"), by.get("042")
        if m and p and (m.pct < 50 or m.passed == "No") and p.pct >= 75:
            want.add(ds.students[sid].name)
    assert {r[0] for r in ans["table"]["rows"]} == want


def test_synonyms_and_typos(data):
    for q, subj in [("average of maths in class 11", "041"), ("avg of mathematics class 11", "041"),
                    ("phy toppers class xii", "042"), ("chemestry toppers class 12", "043")]:
        parsed, ans = ask(data, q)
        assert parsed["query"].get("subject") == subj, q


def test_hinglish(data):
    parsed, ans = ask(data, "physics mein kaun kamzor hai class 11")
    assert parsed["query"]["op"] == "list"
    assert parsed["query"]["conds"][0] == {"level": "low", "subject": "042"}


def test_speech_to_text_slips(data):
    # speech-to-text often writes "weak" as "week"
    parsed, ans = ask(data, "Who is Week in Physics in Class 11?")
    assert parsed["query"]["conds"] == [{"level": "low", "subject": "042"}]


def test_unsupported_question_says_so(data):
    parsed, ans = ask(data, "what is the capital of france")
    assert parsed.get("unknown") and ans is None and parsed["suggestions"]


def test_ambiguous_question_asks_back(data):
    parsed, ans = ask(data, "top 10 students")
    assert parsed.get("clarify") and {o["force"]["cls"] for o in parsed["options"]} == {"VI", "IX", "X", "XI", "XII"}
    parsed2, ans2 = ask(data, "top 10 students", parsed["options"][0]["force"])
    assert ans2 and len(ans2["table"]["rows"]) >= 10


def test_malformed_query_rejected(data):
    ds, res = data
    E = Engine(ds, res)
    from markstudio.query import QueryError
    for bad in [{"op": "drop_table"}, {"op": "top", "n": -3, "scope": {}}, {"op": "top", "subject": "999", "scope": {}},
                {"op": "list", "conds": [{"label": "Genius"}], "scope": {}}, "not a dict"]:
        with pytest.raises(QueryError):
            E.run(bad)


def test_compare_sections(data):
    parsed, ans = ask(data, "compare section A and B in class 11")
    assert parsed["query"]["op"] == "compare" and len(parsed["query"]["groups"]) == 2
    assert len(ans["table"]["rows"]) == 2


def test_word_commands(data):
    ds, res = data
    api.S = api.Session()
    api.S.ds, api.S.res = ds, res
    out = json.loads(api.command("add a chart of physics averages by section"))["data"]
    assert out["item"]["type"] == "chart" and out["item"]["by"] == "section" and out["item"]["filters"]["subject"] == "042"
    assert out["preview"]["chart"]["values"]
    out = json.loads(api.command("highlight students below 40% in red"))["data"]
    assert out["item"] == {"type": "highlight", "op": "<", "value": 40.0, "color": "red", "subject": None, "exam": None}
    out = json.loads(api.command("change the weak threshold to 8"))["data"]
    assert out["item"] == {"type": "setting", "key": "weak_gap", "value": 8}
    json.loads(api.apply_command(json.dumps(out["item"])))
    assert api.S.cfg.rules["weak_gap"] == 8
    json.loads(api.undo())
    assert api.S.cfg.rules["weak_gap"] == 5
    out = json.loads(api.command("add a sheet with the top 10 in class 12"))["data"]
    assert out["item"]["type"] == "query_sheet" and len(out["item"]["rows"]) >= 10
    out = json.loads(api.command("make the logo spin"))["data"]
    assert out.get("error")


def test_workbook_with_word_command_extras(data):
    ds, res = data
    recipe = {"items": [{"type": "chart", "metric": "average", "by": "section", "filters": {"cls": "XI", "subject": "042", "exam": "YE"},
                         "kind": "column", "title": "Physics by section"},
                        {"type": "highlight", "op": "<", "value": 40, "color": "red", "subject": None, "exam": None},
                        {"type": "query_sheet", "sheet": "Top Ten", "title": "Top ten", "explanation": "x",
                         "columns": ["A", "B"], "rows": [["x", 1], ["y", 2.5]]}]}
    import openpyxl
    b = write_workbook(ds, res, recipe=recipe)
    wb = openpyxl.load_workbook(io.BytesIO(b))
    assert "My Charts" in wb.sheetnames and "Top Ten" in wb.sheetnames
