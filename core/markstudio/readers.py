"""Turns any uploaded file into plain grids of text (one grid per sheet/page).

Nothing here decides what a column means - that is layout.py's job. Here we only
make sure every cell is read exactly as the teacher typed it."""
import csv
import io
import re
import datetime
from decimal import Decimal

from .textutil import clean


class Sheet:
    def __init__(self, name, rows, kind="table"):
        self.name = name
        self.rows = rows              # list of lists of strings (ragged allowed)
        self.kind = kind              # 'table' | 'text'
        self.text = None              # full text when the sheet is really plain text

    def width(self):
        return max((len(r) for r in self.rows), default=0)


def cell_text(v):
    """A spreadsheet value as the text the teacher would see/type."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        d = Decimal(repr(v))
        if d == d.to_integral_value():
            return str(int(d))
        return format(d.normalize(), "f")
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%d-%m-%Y")
    return clean(v)


# ------------------------------------------------------------------ text files
def decode_text(data):
    for enc in ("utf-8-sig", "utf-16"):
        try:
            t = data.decode(enc)
            if enc == "utf-16" and not (data[:2] in (b"\xff\xfe", b"\xfe\xff")):
                continue
            return t
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def _sniff_delimiter(lines):
    sample = [l for l in lines if l.strip()][:50]
    best, best_score = None, 0
    for d in [",", "\t", ";", "|"]:
        counts = [l.count(d) for l in sample]
        if not counts or max(counts) == 0:
            continue
        common = max(set(counts), key=counts.count)
        score = counts.count(common) * common if common else 0
        if score > best_score:
            best, best_score = d, score
    return best


def read_text(name, data):
    text = decode_text(data) if isinstance(data, (bytes, bytearray)) else data
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    delim = _sniff_delimiter(lines)
    if delim:
        rows = [[clean(c) for c in r] for r in csv.reader(io.StringIO("\n".join(lines)), delimiter=delim)]
        sh = Sheet(name, rows)
    else:
        # fixed-width text: split on runs of 2+ spaces
        rows = [[clean(c) for c in re.split(r"\s{2,}", l.strip())] if l.strip() else [] for l in lines]
        sh = Sheet(name, rows, kind="text")
    sh.text = text
    return [sh]


# ------------------------------------------------------------------ Excel
def read_xlsx(name, data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        maxr, maxc = ws.max_row or 0, ws.max_column or 0
        grid = [["" for _ in range(maxc)] for _ in range(maxr)]
        for row in ws.iter_rows(min_row=1, max_row=maxr, max_col=maxc):
            for c in row:
                if c.value is not None:
                    grid[c.row - 1][c.column - 1] = cell_text(c.value)
        # merged cells: copy the top-left value into the whole merged area
        merged = []
        for mr in ws.merged_cells.ranges:
            v = grid[mr.min_row - 1][mr.min_col - 1] if mr.min_row - 1 < maxr and mr.min_col - 1 < maxc else ""
            merged.append((mr.min_row - 1, mr.max_row - 1, mr.min_col - 1, mr.max_col - 1))
            for r in range(mr.min_row - 1, min(mr.max_row, maxr)):
                for c in range(mr.min_col - 1, min(mr.max_col, maxc)):
                    grid[r][c] = v
        sh = Sheet(ws.title, _trim(grid))
        sh.merged = merged
        sheets.append(sh)
    return sheets


def read_xls(name, data):
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    except Exception:
        book = xlrd.open_workbook(file_contents=data)
    sheets = []
    for sh in book.sheets():
        if getattr(sh, "visibility", 0) != 0:
            continue
        grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                    except Exception:
                        v = cell.value
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    v = "#ERROR"
                else:
                    v = cell.value
                row.append(cell_text(v))
            grid.append(row)
        for (r0, r1, c0, c1) in getattr(sh, "merged_cells", []):
            v = grid[r0][c0] if r0 < len(grid) and c0 < len(grid[r0]) else ""
            for r in range(r0, min(r1, len(grid))):
                for c in range(c0, min(c1, len(grid[r]))):
                    grid[r][c] = v
        sheets.append(Sheet(sh.name, _trim(grid)))
    return sheets


def _trim(grid):
    """Drop empty trailing rows/columns."""
    while grid and not any(x for x in grid[-1]):
        grid.pop()
    w = 0
    for r in grid:
        for i in range(len(r) - 1, -1, -1):
            if r[i]:
                w = max(w, i + 1)
                break
    return [r[:w] for r in grid]


def read_file(name, data):
    """Pick the reader from the file's content (not just its name)."""
    head = bytes(data[:8])
    low = name.lower()
    if head.startswith(b"PK"):
        return read_xlsx(name, data)
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return read_xls(name, data)
    if low.endswith((".xlsx", ".xlsm", ".xls")):
        raise ValueError("This Excel file could not be opened. Is it password-protected or damaged?")
    return read_text(name, data)


# ------------------------------------------------------------------ PDF / scans
def lines_from_words(words, char_w=None):
    """Words with positions (from a PDF page or a scan) -> text lines that keep the
    gaps between columns, so the same fixed-width parsers work on PDFs."""
    if not words:
        return []
    heights = sorted(max(1.0, w["y1"] - w["y0"]) for w in words)
    h = heights[len(heights) // 2]
    if char_w is None:
        widths = sorted((w["x1"] - w["x0"]) / max(1, len(w["text"])) for w in words if w["text"].strip())
        char_w = widths[len(widths) // 2] if widths else 5.0
    rows = []
    for w in sorted(words, key=lambda w: ((w["y0"] + w["y1"]) / 2, w["x0"])):
        cy = (w["y0"] + w["y1"]) / 2
        if rows and abs(rows[-1]["cy"] - cy) <= h * 0.5:
            rows[-1]["words"].append(w)
            n = len(rows[-1]["words"])
            rows[-1]["cy"] = (rows[-1]["cy"] * (n - 1) + cy) / n
        else:
            rows.append({"cy": cy, "words": [w]})
    out = []
    for r in rows:
        line = ""
        for w in sorted(r["words"], key=lambda w: w["x0"]):
            col = int(round(w["x0"] / char_w))
            if len(line) < col:
                line += " " * (col - len(line))
            elif line and not line.endswith(" "):
                line += " "
            line += w["text"]
        out.append(line.rstrip())
    return out


def grid_from_words(words, join_words=True):
    """Words with positions -> a table grid.
    Rows: words whose vertical centres line up (a small tilt in a scan is allowed).
    Columns: found from the rows that are mostly numbers (the marks), because headings
    are often aligned differently (left vs right); every heading is then put in the
    nearest column. Nothing is dropped: every word ends up in some cell."""
    if not words:
        return []
    heights = sorted(max(1.0, w["y1"] - w["y0"]) for w in words)
    h = heights[len(heights) // 2]
    # rows (allow for a tilted scan by following each row's own slope)
    rows = []
    for w in sorted(words, key=lambda w: (w["x0"])):
        cy = (w["y0"] + w["y1"]) / 2
        best = None
        for r in rows:
            last = r["words"][-1]
            ly = (last["y0"] + last["y1"]) / 2
            if abs(ly - cy) <= h * 0.55 and (best is None or abs(ly - cy) < best[0]):
                best = (abs(ly - cy), r)
        if best:
            best[1]["words"].append(w)
        else:
            rows.append({"words": [w]})
    for r in rows:
        r["cy"] = sum((w["y0"] + w["y1"]) / 2 for w in r["words"]) / len(r["words"])
    rows.sort(key=lambda r: r["cy"])
    # cells: join words that are close together (scans give single words; PDFs give whole cells)
    gap_join = h * 0.9
    cells_by_row = []
    for r in rows:
        ws = sorted(r["words"], key=lambda w: w["x0"])
        cells = []
        for w in ws:
            if join_words and cells and w["x0"] - cells[-1]["x1"] <= gap_join:
                cells[-1]["text"] += " " + w["text"]
                cells[-1]["x1"] = max(cells[-1]["x1"], w["x1"])
            else:
                cells.append({"text": w["text"], "x0": w["x0"], "x1": w["x1"]})
        cells_by_row.append(cells)

    def numeric_share(cells):
        n = sum(1 for c in cells if re.match(r"^[\d.]+$", c["text"].strip()) or c["text"].strip().upper() in ("AB", "ML", "-"))
        return n / max(1, len(cells))
    most = max(len(c) for c in cells_by_row)
    body = [c for c in cells_by_row if len(c) >= max(2, most * 0.5) and numeric_share(c) >= 0.4]
    if not body:
        body = [c for c in cells_by_row if len(c) >= max(2, most * 0.6)] or cells_by_row
    spans = sorted((c["x0"], c["x1"]) for cells in body for c in cells)
    cols = []
    for x0, x1 in spans:
        if cols and x0 <= cols[-1][1] + h * 0.3:
            cols[-1][1] = max(cols[-1][1], x1)
        else:
            cols.append([x0, x1])
    grid = []
    for cells in cells_by_row:
        row = [""] * len(cols)
        for c in cells:
            row[_column_for(c, cols)] = (row[_column_for(c, cols)] + " " + c["text"]).strip()
        grid.append(row)
    return grid


def _column_for(c, cols):
    """Which column a cell belongs to: the column it overlaps most; if it overlaps none
    (a left-aligned heading above right-aligned numbers), the first column to its right."""
    best, best_ov = None, 0
    for i, (a, b) in enumerate(cols):
        ov = min(b, c["x1"]) - max(a, c["x0"])
        if ov > best_ov:
            best, best_ov = i, ov
    if best is not None:
        return best
    right = [i for i, (a, b) in enumerate(cols) if b >= c["x0"]]
    if right:
        return min(right, key=lambda i: cols[i][1])
    mid = (c["x0"] + c["x1"]) / 2
    return min(range(len(cols)), key=lambda i: min(abs(mid - cols[i][0]), abs(mid - cols[i][1])))


def read_pages(name, pages):
    """pages: [{'page': 1, 'words': [{'text','x0','x1','y0','y1'}], 'ocr': bool}]
    Returns sheets (one per page) plus the whole document as fixed-width text."""
    sheets = []
    all_lines = []
    for p in pages:
        grid = grid_from_words(p["words"], join_words=bool(p.get("ocr")))
        sh = Sheet(f"Page {p['page']}", grid)
        sh.ocr = bool(p.get("ocr"))
        sheets.append(sh)
        all_lines.extend(lines_from_words(p["words"]))
    text = "\n".join(all_lines)
    for sh in sheets:
        sh.text = text
    return sheets, text
