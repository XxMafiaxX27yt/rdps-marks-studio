"""Everything the website calls. Each function takes simple values (text,
numbers, bytes, JSON text) and returns JSON-friendly data, so it works the same
inside the browser (Pyodide) and in the tests.

All work happens on this computer. Nothing here talks to the internet."""
import json
import copy
import datetime
import traceback

from .config import Config, default_settings
from .build import Project, Builder
from . import calc, verify, views, gazette, board, snapshot
from .nlu import Understander
from .query import Engine, QueryError
from .commands import parse_command
from .views import to_json

VERSION = "2.0.0"


class Session:
    def __init__(self, settings=None):
        self.cfg = Config(settings)
        self.project = Project(self.cfg)
        self.ds = None
        self.res = None
        self.verify = None
        self.recipe = {"items": []}
        self.undo_stack = []
        self.restored = None        # file name when a workbook was loaded back
        self.boards = []
        self.board_texts = []
        self.teachers = []          # typed on the website (the browser keeps them)


S = Session()


def _ok(data):
    return json.dumps({"ok": True, "data": to_json(data)}, default=str)


def _err(msg, detail=""):
    return json.dumps({"ok": False, "error": msg, "detail": detail})


def safe(fn):
    def wrap(*a, **kw):
        try:
            return _ok(fn(*a, **kw))
        except QueryError as ex:
            return _err(str(ex))
        except verify.VerificationError as ex:
            return _err("Safety check failed - the two independent calculations disagree, so no results were made.", str(ex))
        except Exception as ex:  # shown to the user in plain words, details for us
            return _err(str(ex) or ex.__class__.__name__, traceback.format_exc())
    wrap.__name__ = fn.__name__
    return wrap


# ================================================================== settings
@safe
def init(settings_json=""):
    global S
    settings = json.loads(settings_json) if settings_json else None
    S = Session(settings)
    return {"version": VERSION, "settings": S.cfg.settings}


@safe
def get_settings():
    return S.cfg.settings


@safe
def default_settings_json():
    return default_settings()


@safe
def set_settings(settings_json):
    new = Config(json.loads(settings_json))
    old_project = S.project
    S.cfg = new
    p = Project(new)
    p.remembered = old_project.remembered
    p.corrections = old_project.corrections
    for src in old_project.sources:
        p.sources.append(src)
        from .layout import detect, signature
        for sh in src.sheets:
            old = old_project.layouts.get((src.id, sh.name))
            if old and old.get("kind") == "info":
                p.layouts[(src.id, sh.name)] = old
                continue
            lay = detect(sh, src.name, new, p.vocab)
            lay["signature"] = signature(sh, lay)
            if lay["signature"] in p.remembered:
                p._apply_saved(lay, p.remembered[lay["signature"]])
            p.layouts[(src.id, sh.name)] = lay
    p._n = old_project._n
    p.gazettes = old_project.gazettes
    p.file_teachers = old_project.file_teachers
    p.class_lists, p._cl = old_project.class_lists, old_project._cl
    S.project = p
    if S.restored and S.ds is not None:
        # a workbook that was opened again has no files to read: keep its data, work it out with the new rules
        S.ds.cfg = new
        S.res = None
    else:
        S.ds = S.res = None
    return {"settings": new.settings}


@safe
def remember_json():
    """Saved column choices (the website keeps them in the browser between visits)."""
    return S.project.remembered


@safe
def load_remembered(js):
    S.project.remembered.update(json.loads(js))
    return {"count": len(S.project.remembered)}


# ================================================================== files
def _leave_restored():
    """New marks files after a saved workbook was opened: start from the files (the workbook's data is dropped)."""
    if S.restored:
        S.restored = None
        S.ds = S.res = None
        S.recipe = {"items": []}


@safe
def add_file(name, data):
    data = bytes(data)
    if data[:2] == b"PK" and name.lower().endswith((".xlsx", ".xlsm")):
        try:
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
            if snapshot.is_ours(wb.sheetnames):
                return {"kind": "our_workbook", "name": name}
        except Exception:
            pass
    info = S.project.add_file(name, data)
    if info["kind"] == "gazette":
        _refresh_boards()
    else:
        _leave_restored()
    return info


@safe
def add_pages(name, pages_json):
    info = S.project.add_pages(name, json.loads(pages_json))
    if info["kind"] != "gazette":
        _leave_restored()
    if info["kind"] == "gazette":
        _refresh_boards()
    return info


@safe
def add_sample(mistakes=False):
    from .sample import sample_files
    _leave_restored()
    out = []
    for name, data in sample_files(S.cfg, mistakes=bool(mistakes)):
        out.append(S.project.add_file(name, data))
    return out


@safe
def add_any(name, data, prefer="", hint=""):
    """Any file from any page: sorted into board results (XII / X), a class list or school marks."""
    data = bytes(data)
    if data[:2] == b"PK" and name.lower().endswith((".xlsx", ".xlsm")) and not prefer:
        try:
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
            if snapshot.is_ours(wb.sheetnames):
                return {"kind": "our_workbook", "name": name}
        except Exception:
            pass
    return _after_any(S.project.add_any(name, data, prefer, hint))


@safe
def add_any_pages(name, pages_json, prefer="", hint=""):
    return _after_any(S.project.add_any_pages(name, json.loads(pages_json), prefer, hint))


def _after_any(info):
    if info["kind"] == "board":
        _refresh_boards()
    elif info["kind"] == "school":
        _leave_restored()
    if info["kind"] == "class_list":
        info["class_lists"] = _class_lists_out()
    return info


@safe
def add_board_file(name, data):
    """From the Board page: CBSE's gazette or the school's own board-results table."""
    info = S.project.add_board_file(name, bytes(data))
    _refresh_boards()
    return info


@safe
def add_board_pages(name, pages_json):
    info = S.project.add_board_pages(name, json.loads(pages_json))
    _refresh_boards()
    return info


# ================================================================== class lists
def _class_lists_out():
    return [{"id": c["id"], "file": c["file"], "sheet": c["sheet"], "cls": c["cls"], "section": c["section"],
             "needs": c["needs"], "count": len(c["students"]),
             "sections": sorted({f"{s['cls']}-{s['section']}" if s["section"] else s["cls"] for s in c["students"]})}
            for c in S.project.class_lists]


@safe
def add_class_list(name, data):
    S.project.add_class_list(name, bytes(data))
    return _class_lists_out()


@safe
def add_class_list_pages(name, pages_json):
    S.project.add_class_list_pages(name, json.loads(pages_json))
    return _class_lists_out()


@safe
def class_lists():
    return _class_lists_out()


@safe
def set_class_list(list_id, cls, section):
    from .classlist import set_class
    for c in S.project.class_lists:
        if c["id"] == list_id:
            set_class(c, cls, section)
    return _class_lists_out()


@safe
def remove_class_list(list_id):
    S.project.class_lists = [c for c in S.project.class_lists if c["id"] != list_id]
    return _class_lists_out()


@safe
def add_sample_gazette():
    from .sample import sample_gazette
    info = S.project.add_file("Sample gazette (made-up).txt", sample_gazette(S.cfg).encode())
    _refresh_boards()
    return info


@safe
def remove_file(source_id):
    if source_id.startswith("G"):
        i = int(source_id[1:])
        del S.project.gazettes[i]
        _refresh_boards()
    else:
        S.project.remove_source(source_id)
    S.ds = S.res = None
    return state_dict()


@safe
def clear_all():
    global S
    settings = S.cfg.settings
    remembered = S.project.remembered
    teachers = S.teachers
    S = Session(settings)
    S.project.remembered = remembered
    S.teachers = teachers
    return {"cleared": True}


def state_dict():
    p = S.project
    files = []
    for src in p.sources:
        sheets = []
        for sh in src.sheets:
            lay = p.layouts[(src.id, sh.name)]
            preview = []
            hdr = set(lay.get("header_rows", []))
            for r in range(min(len(sh.rows), lay.get("data_start", 0) + 6)):
                preview.append({"row": r + 1, "header": r in hdr, "cells": sh.rows[r][:60]})
            sheets.append({
                "name": sh.name, "kind": lay.get("kind", "marks"), "skip": lay["skip"], "note": lay.get("note", ""),
                "exam": lay.get("exam"), "exam_from": lay.get("exam_from", ""), "cls": lay.get("cls"),
                "cls_from": lay.get("cls_from", ""), "section": lay.get("section", ""),
                "needs_exam": lay.get("needs_exam", False), "needs_class": lay.get("needs_class", False),
                "format": lay.get("format", "wide"), "rows": len(sh.rows), "data_start": lay.get("data_start", 0) + 1,
                "remembered": lay.get("remembered", False), "title": lay.get("title", ""),
                "columns": [{k: c.get(k) for k in ("col", "header", "role", "subject", "component", "exam", "sure", "why", "sample")}
                            for c in lay["columns"]],
                "preview": preview, "ocr": getattr(sh, "ocr", False),
            })
        files.append({"id": src.id, "name": src.name, "kind": src.kind, "sheets": sheets})
    gz = [{"id": f"G{i}", "name": g.file, "year": g.year, "cls": g.cls, "students": len(g.students),
           "school": g.school_name, "errors": sum(1 for x in g.issues if x.severity == "error"),
           "what": getattr(g, "kind", "gazette")}
          for i, g in enumerate(p.gazettes)]
    cls_lists = [{"id": c["id"], "name": c["file"], "count": len(c["students"]), "needs": c["needs"],
                  "sections": sorted({f"{s['cls']}-{s['section']}" if s["section"] else s["cls"] for s in c["students"]
                                      if s["cls"]})} for c in p.class_lists]
    return {"files": files, "gazettes": gz, "class_lists": cls_lists, "pending": p.pending(), "corrections": p.corrections,
            "restored": S.restored, "has_results": S.res is not None}


@safe
def state():
    return state_dict()


@safe
def set_sheet(source_id, sheet, js):
    S.project.set_sheet(source_id, sheet, **json.loads(js))
    S.ds = S.res = None
    return state_dict()


@safe
def set_column(source_id, sheet, col, js):
    d = json.loads(js)
    S.project.set_column(source_id, sheet, int(col), d.get("role", "ignore"), d.get("subject"), d.get("component"), d.get("exam"))
    S.ds = S.res = None
    return state_dict()


@safe
def confirm_sheet(source_id, sheet):
    S.project.confirm_sheet(source_id, sheet)
    return state_dict()


@safe
def swap_columns(source_id, sheet, a, b):
    S.project.swap_columns(source_id, sheet, int(a), int(b))
    S.ds = S.res = None
    return state_dict()


@safe
def correct(key, value):
    S.project.correct(key, None if value is None else str(value))
    S.ds = S.res = None
    return {"key": key, "value": value}


# ================================================================== results
def _compute():
    S.res = calc.compute(S.ds)
    S.verify = verify.check(S.ds, S.res)


@safe
def analyze():
    if not S.restored:
        pend = S.project.pending()
        if pend:
            return {"blocked": True, "pending": pend}
        S.ds = S.project.build()
    _compute()
    return summary_dict()


def summary_dict():
    ds, res, cfg = S.ds, S.res, S.cfg
    classes = []
    for c in ds.classes():
        classes.append({"cls": c, "label": cfg.class_label(c),
                        "combos": sorted({s.combo for s in ds.students.values() if s.cls == c and s.combo}),
                        "exams": [{"code": e, "name": cfg.exam_name(e)} for e in ds.exams_for(c)],
                        "sections": ds.sections_for(c),
                        "subjects": [{"code": s, "name": cfg.subject_name(s), "short": cfg.subject_short(s)} for s in ds.subjects_for(c)],
                        "students": sum(1 for s in ds.students.values() if s.cls == c)})
    sev = {"error": 0, "warning": 1, "info": 2}
    issues = sorted(ds.issues, key=lambda i: (sev.get(i.severity, 3), i.code))
    return {
        "students": len(ds.students), "records": len(ds.marks), "results": len(res.rows),
        "classes": classes, "verify": S.verify,
        "issues": [i.to_dict(ds.students) for i in issues],
        "counts": {k: sum(1 for i in issues if i.severity == k) for k in sev},
        "student_list": [{"sid": s.sid, "label": views.student_labels(ds)[s.sid], "name": s.name, "cls": s.cls,
                          "section": s.section, "roll": s.roll, "group": s.group} for s in ds.ordered_students()],
        "restored": S.restored, "recipe": S.recipe,
        "suggestions": Understander(ds, res, S.boards, cfg).suggestions(),
    }


@safe
def summary():
    if S.res is None:
        raise ValueError("No results yet.")
    return summary_dict()


@safe
def dashboard(cls, exam, section="All"):
    v = views.dashboard(S.res, cls, exam, section)
    v["by_subject"] = [{"code": s, "name": S.cfg.subject_short(s), "full": S.cfg.subject_name(s), "avg": a} for s, a in v["by_subject"]]
    return v


@safe
def subject_analysis(cls, exam):
    rows = views.subject_analysis(S.res, cls, exam)
    return {"rows": rows, "practical": [s for s in views.all_subjects(S.res) if S.cfg.has_practical(s)]}


@safe
def subject_trend(cls, subject):
    return [{"exam": e, "name": S.cfg.exam_name(e), "avg": a, "pass": p} for e, a, p in views.subject_trend(S.res, cls, subject)]


@safe
def report_card(sid, exam):
    return views.report_card(S.res, sid, exam)


@safe
def report_cards(cls, exam, section="All"):
    out = []
    for st in S.ds.ordered_students():
        if st.cls != cls or (section not in ("All", "", None) and st.section != section):
            continue
        if (st.sid, exam) in S.res.exam_results:
            out.append(views.report_card(S.res, st.sid, exam))
    return {"cards": out, "school": S.cfg.rules["school_name"], "place": S.cfg.rules["school_place"],
            "year": S.cfg.rules["academic_year"]}


@safe
def progress(cls):
    rows = views.progress_rows(S.res, cls)
    return {"exams": [{"code": e, "name": S.cfg.exam_name(e)} for e in S.ds.exams_for(cls)], "rows": rows}


@safe
def matrix(cls, exam, section="All", combo=""):
    """Marks table for the preview (one row per student). Only the subjects of the class (or the chosen
    stream); each student's language / elective shares one column."""
    from .excel.school import _slot_plan
    ds, res, cfg = S.ds, S.res, S.cfg
    studs = [st for st in ds.ordered_students() if st.cls == cls and (section in ("All", "", None) or st.section == section)
             and (not combo or st.combo == combo) and (st.sid, exam) in res.exam_results]
    subjects = cfg.order_subjects({r.subject for st in studs for r in res.exam_results[(st.sid, exam)].rows})
    core, slots, slot_of = _slot_plan(cfg, cls, studs, subjects)
    rows = []
    for st in studs:
        er = res.exam_results[(st.sid, exam)]
        by = {r.subject: r for r in er.rows}

        def cell(r):
            if r is None:
                return None
            return {"pct": r.pct, "label": r.label, "status": r.status, "th": r.th if r.th is not None else r.th_raw,
                    "pr": (r.pr if r.pr is not None else r.pr_raw) if r.pr_max else None, "max": r.max,
                    "obt": r.obtained, "reason": r.reason, "subject": cfg.subject_name(r.subject)}
        cells = [cell(by.get(s)) for s in core]
        mine = slot_of.get(st.sid, {})
        for i in range(len(slots)):
            sub = next((x for x, k in mine.items() if k == i), None)
            c = cell(by.get(sub)) if sub else None
            if c:
                c["slot_subject"] = cfg.subject_short(sub)
            cells.append(c)
        rows.append({"sid": st.sid, "name": st.name, "roll": st.roll, "section": st.section, "group": st.group,
                     "combo": st.combo, "cells": cells, "overall": er.overall, "result": er.result, "rank": er.rank})
    cols = [{"code": s, "short": cfg.subject_short(s), "name": cfg.subject_name(s)} for s in core] +            [{"code": "", "short": sl["short"], "name": sl["title"] + ": " + ", ".join(cfg.subject_name(x) for x in
                                                                                    cfg.order_subjects(sl["subjects"]))}
            for sl in slots]
    combos = sorted({st.combo for st in ds.students.values() if st.cls == cls and st.combo})
    return {"subjects": cols, "rows": rows, "combos": combos}


# ================================================================== downloads
def build_workbook_bytes(macros=True):
    """Fast path for the browser: returns raw bytes (no JSON) or raises."""
    from .excel.school import write_workbook
    if S.res is None:
        raise ValueError("Make the results first.")
    return write_workbook(S.ds, S.res, recipe=S.recipe, vba=bool(macros), teachers=all_teachers())


def workbook_name(macros=True):
    stamp = datetime.datetime.now().strftime("%Y-%m-%d")
    return f"RDPS Results {S.cfg.rules['academic_year']} ({stamp}).{'xlsm' if macros else 'xlsx'}"


def template_bytes(cls, prefill=False):
    from .excel.template import template_bytes as tb
    students = None
    if prefill and S.ds:
        students = [{"adm_no": s.adm_no, "roll": s.roll, "board_roll": s.board_roll, "name": s.name, "father": s.father,
                     "mother": s.mother, "section": s.section, "group": s.group, "subjects": s.subjects}
                    for s in S.ds.ordered_students() if s.cls == cls]
    teachers = [t for t in all_teachers() if any(c == cls for c, _ in t["pairs"])]
    return tb(S.cfg, cls, students=students, teachers=teachers)


# ================================================================== teachers
def all_teachers():
    """Teachers typed on the website + teachers read from 'Teachers' sheets. The same name + subject is one
    teacher: a typed entry wins; the same teacher in several files (one per class) gets all those classes."""
    from .teachers import normalise
    from .textutil import norm_name
    typed = normalise(S.teachers, S.cfg, S.project.vocab)
    typed_keys = {(norm_name(t["name"]), t["subject"]) for t in typed}
    out = list(typed)
    merged = {}
    for t in normalise(S.project.file_teachers, S.cfg, S.project.vocab):
        k = (norm_name(t["name"]), t["subject"])
        if k in typed_keys:
            continue
        if k in merged:
            m = merged[k]
            for p in t["pairs"]:
                if p not in m["pairs"]:
                    m["pairs"].append(p)
                    m["classes"].append(f"{p[0]}-{p[1]}" if p[1] else p[0])
            continue
        t["id"] = f"F{len(out) + 1}"
        merged[k] = t
        out.append(t)
    return out


@safe
def set_teachers(js):
    S.teachers = json.loads(js) if js else []
    return teachers_list_inner()


def teachers_list_inner():
    from .teachers import normalise
    return {"typed": normalise(S.teachers, S.cfg, S.project.vocab),
            "from_files": normalise(S.project.file_teachers, S.cfg, S.project.vocab)}


@safe
def remove_file_teacher(name, subject):
    """Remove a teacher read from an uploaded 'Teachers' sheet (completely - from the list, the analysis and
    the Excel file). name + subject as shown in the list."""
    from .teachers import normalise
    from .textutil import norm_name
    keep = []
    for t, n in zip(S.project.file_teachers, normalise(S.project.file_teachers, S.cfg, S.project.vocab)):
        if norm_name(n["name"]) == norm_name(name) and (n["subject"] or n["subject_text"]) == subject:
            continue
        keep.append(t)
    S.project.file_teachers = keep
    return teachers_list_inner()


@safe
def remove_all_teachers():
    S.teachers = []
    S.project.file_teachers = []
    return teachers_list_inner()


@safe
def teachers_list():
    return teachers_list_inner()


@safe
def teacher_analysis():
    from .teachers import analyse
    if S.res is None:
        raise ValueError("Make the results first.")
    ts = all_teachers()
    out = analyse(S.res, ts)
    for c in out["cards"]:
        c.pop("pairs", None)
    out["count"] = len(ts)
    return out


# ================================================================== boards
def _refresh_boards():
    S.boards = [board.analyse(g, S.cfg) for g in S.project.gazettes]


def _board_label(an):
    m = an["meta"]
    lab = f"Class {m.get('cls') or 'XII'} · {m.get('year') or ''}".strip(" ·")
    if (m.get("exam_type") or "MAIN").upper() not in ("MAIN", ""):
        lab += f" · {m['exam_type'].title()}"
    return lab


def _class_boards(cls=""):
    """The board results of one class (default: the class of the file added last), oldest year first."""
    if not S.boards:
        return []
    cls = cls or (S.boards[-1]["meta"].get("cls") or "")
    same = [a for a in S.boards if (a["meta"].get("cls") or "") == cls]
    return sorted(same, key=lambda a: a["meta"]["year"])


@safe
def board_summary():
    if not S.project.gazettes:
        raise ValueError("Upload a CBSE gazette (.TXT) first.")
    _refresh_boards()
    out = []
    for i, an in enumerate(S.boards):
        d = {k: v for k, v in an.items() if k not in ("rows", "ranked", "comp_list", "others_list", "issues",
                                                      "all_rounders")}
        bcls = an["meta"].get("cls") or ""
        d["key"], d["label"] = f"{i}", _board_label(an)
        d["ranked"] = [{"rank": r["rank"], "name": r["s"].name, "roll": r["s"].roll, "best5": r["best5"],
                        "pct_all": r["pct_all"], "result": r["s"].result, "stream": r["stream"]} for r in an["ranked"][:25]]
        d["comp_list"] = [{"name": s.name, "roll": s.roll, "comp": [S.cfg.board_name(c, bcls) for c in s.comp],
                           "marks": [x["marks"] for x in s.subjects if x["code"] in s.comp]} for s in an["comp_list"]]
        d["others_list"] = [{"name": s.name, "roll": s.roll, "result": s.result} for s in an["others_list"]]
        d["all_rounders"] = [r["s"].name for r in an["all_rounders"]]
        d["issues"] = [x.to_dict() for x in an["issues"]]
        d["highlights"] = board.highlights(an, S.cfg)
        d["students"] = [{"roll": r["s"].roll, "name": r["s"].name, "gender": r["s"].gender, "result": r["s"].result,
                          "best5": r["best5"], "pct_all": r["pct_all"], "rank": r["rank"], "stream": r["stream"],
                          "combo": r["combo"], "section": getattr(r["s"], "section", "") or "",
                          "best5_subjects": r["best5_subjects"], "comp": list(r["s"].comp),
                          "subjects": [{"code": x["code"], "name": S.cfg.board_short(x["code"], bcls), "marks": x["marks"],
                                        "grade": x["grade"]} for x in r["s"].subjects]} for r in an["rows"]]
        same = _class_boards(bcls)
        d["years"] = board.compare_years(same) if len({a["meta"]["year"] for a in same}) > 1 else None
        link = board.school_vs_board(an, S.ds, S.res, S.cfg) if S.res else None
        d["link"] = {k: v for k, v in link.items() if k != "links"} if link else None
        d["teachers"] = _board_teachers(an)
        out.append(d)
    for d in out:                       # two files of the same class and year: say which file is which
        if sum(1 for x in out if x["label"].split(" · ")[:2] == d["label"].split(" · ")[:2]) > 1:
            d["label"] += " · " + d["meta"]["file"]
    last = out[-1]
    return {"boards": out, "years": last["years"], "link": last["link"], "teachers": last["teachers"]}


def _book_set(cls="", key=""):
    """The board results for one workbook: the file picked on the Board page (key), or the latest of a class,
    plus the other years of the same class to compare with. The main one is last."""
    k = str(key or "")
    if k.isdigit() and int(k) < len(S.boards):
        main = S.boards[int(k)]
        others = [a for a in _class_boards(main["meta"].get("cls") or "") if a["meta"]["year"] != main["meta"]["year"]]
        return others + [main]
    return _class_boards(cls)


def board_workbook_bytes(cls="", key=""):
    """The board workbook of one class (Class XII or X); several years of that class are compared."""
    from .excel.board_book import write_board_workbook
    _refresh_boards()
    an = _book_set(cls, key)
    if not an:
        raise ValueError(f"No board results for Class {cls}.")
    years = board.compare_years(an) if len({a["meta"]["year"] for a in an}) > 1 else None
    link = board.school_vs_board(an[-1], S.ds, S.res, S.cfg) if S.res else None
    files = {a["meta"]["file"] for a in an}
    texts = [(g.file, getattr(g, "raw_text", "")) for g in S.project.gazettes if g.file in files]
    return write_board_workbook(an, S.cfg, years=years, link=link, raw_texts=texts,
                                teachers=_board_teachers(an[-1]), highlights=board.highlights(an[-1], S.cfg))


def _board_teachers(an=None):
    from .teachers import analyse_board
    an = an or S.boards[-1]
    out = analyse_board(an, all_teachers(), S.cfg, S.ds, S.project.class_lists)
    return out


@safe
def board_teachers(cls=""):
    """Teachers on the latest board results of a class (Teachers page, Board results tab)."""
    if not S.project.gazettes:
        raise ValueError("Upload the board results on the Board page first.")
    _refresh_boards()
    return _board_teachers((_class_boards(cls) or S.boards)[-1])


@safe
def add_sample_board_demo():
    """Made-up board results (school's own table layout) + made-up class lists + made-up teachers."""
    from .sample import sample_board_table, sample_class_lists, sample_board_teachers
    name, text = sample_board_table(S.cfg)
    S.project.add_board_file(name, text.encode())
    for fname, data in sample_class_lists(S.cfg):
        S.project.add_class_list(fname, data)
    have = {(t["name"], t["subject"]) for t in S.project.file_teachers}
    for t in sample_board_teachers():
        if (t["name"], t["subject"]) not in have:
            S.project.file_teachers.append(dict(t, source="made-up sample"))
    _refresh_boards()
    return {"class_lists": _class_lists_out()}


def board_workbook_name(cls="", key=""):
    an = _book_set(cls, key)
    if not an:
        return "RDPS Board Results.xlsx"
    m = an[-1]["meta"]
    return f"RDPS Board Results {m['year']} Class {m.get('cls') or 'XII'}.xlsx"


# ================================================================== ask + commands
@safe
def ask(question, hints_json="", force_json=""):
    hints = json.loads(hints_json) if hints_json else None
    force = json.loads(force_json) if force_json else None
    U = Understander(S.ds, S.res, S.boards, S.cfg, teachers=all_teachers())
    parsed = U.parse(question, hints, force)
    if "query" not in parsed:
        return parsed
    E = Engine(S.ds, S.res, S.boards, teachers=all_teachers(), cfg=S.cfg, class_lists=S.project.class_lists)
    try:
        ans = E.run(parsed["query"])
    except QueryError as ex:
        return {"unknown": True, "message": str(ex), "suggestions": U.suggestions()}
    ans["understood"] = parsed["understood"]
    ans["matched_by"] = parsed.get("matched_by", "")
    ans["query"] = parsed["query"]
    return ans


@safe
def run_query(query_json):
    return Engine(S.ds, S.res, S.boards, teachers=all_teachers(), cfg=S.cfg,
                  class_lists=S.project.class_lists).run(json.loads(query_json))


@safe
def command(text):
    if S.res is None:
        raise ValueError("Make the results first.")
    out = parse_command(text, S.ds, S.res, S.recipe, S.cfg)
    if out.get("sheet_question"):
        U = Understander(S.ds, S.res, S.boards, S.cfg, teachers=all_teachers())
        parsed = U.parse(out["sheet_question"])
        if "query" not in parsed:
            return {"error": "I couldn't understand what the sheet should show: " + parsed.get("message", parsed.get("question", "")),
                    "suggestions": parsed.get("suggestions", [])}
        ans = Engine(S.ds, S.res, S.boards).run(parsed["query"])
        if not ans.get("table"):
            return {"error": "That answer has no table to put on a sheet."}
        name = _sheet_name(out["sheet_question"])
        item = {"type": "query_sheet", "sheet": name, "title": out["sheet_question"].capitalize(),
                "explanation": ans["text"] + " " + ans.get("how", ""), "columns": ans["table"]["columns"],
                "rows": ans["table"]["rows"], "query": parsed["query"]}
        return {"item": item, "description": f"Add a sheet “{name}” with {len(item['rows'])} rows: {ans['text']}",
                "preview": {"table": ans["table"]}}
    if out.get("item", {}).get("type") == "chart":
        out["preview"] = _chart_preview(out["item"])
    return out


def _sheet_name(q):
    import re
    base = re.sub(r"[\[\]:*?/\\']", "", q).strip().title()[:28] or "Answer"
    existing = {it.get("sheet") for it in S.recipe["items"]}
    name, k = base, 2
    while name in existing:
        name = f"{base[:25]} {k}"
        k += 1
    return name


def _chart_preview(item):
    from .excel.extras import categories, py_value

    class _S:
        pass
    sw = _S()
    sw.cfg, sw.ds, sw.res = S.cfg, S.ds, S.res
    sw.subjects, sw.sections = views.all_subjects(S.res), views.all_sections(S.res)
    sw.exams, sw.classes = views.all_exams(S.res), S.ds.classes()
    cats = categories(sw, item["by"], item["filters"])
    labels, values = [], []
    for code, name, short in cats:
        v = py_value(sw, item["metric"], item["by"], code, item["filters"])
        if v is not None:
            labels.append(short)
            values.append(float(round(v, 2)))
    return {"chart": {"type": item.get("kind", "column"), "labels": labels, "values": values, "label": item["title"]}}


@safe
def apply_command(item_json):
    item = json.loads(item_json)
    t = item.get("type")
    if t == "undo":
        return undo_inner()
    if t == "setting":
        key, val = item["key"], item["value"]
        old = S.cfg.rules.get(key)
        S.cfg.rules[key] = val
        S.cfg.settings["rules"][key] = val
        S.undo_stack.append(("setting", key, old))
        if S.ds is not None:
            _compute()
        return {"applied": item, "recipe": S.recipe, "settings": S.cfg.settings}
    if t in ("chart", "highlight", "query_sheet", "hide_sheet"):
        S.recipe["items"].append(item)
        S.undo_stack.append(("item", len(S.recipe["items"]) - 1))
        return {"applied": item, "recipe": S.recipe}
    raise ValueError("Unknown change.")


def undo_inner():
    if not S.undo_stack:
        return {"undone": None, "message": "Nothing to undo.", "recipe": S.recipe}
    kind = S.undo_stack.pop()
    if kind[0] == "item":
        it = S.recipe["items"].pop(kind[1])
        return {"undone": it, "message": "Removed: " + it.get("title", it.get("sheet", it["type"])), "recipe": S.recipe}
    _, key, old = kind
    S.cfg.rules[key] = old
    S.cfg.settings["rules"][key] = old
    if S.ds is not None:
        _compute()
    return {"undone": {"type": "setting", "key": key}, "message": f"Setting {key} back to {old}.", "recipe": S.recipe,
            "settings": S.cfg.settings}


@safe
def undo():
    return undo_inner()


@safe
def recipe():
    return S.recipe


@safe
def remove_recipe_item(index):
    it = S.recipe["items"].pop(int(index))
    S.undo_stack = [u for u in S.undo_stack if u[0] != "item"]
    return {"removed": it, "recipe": S.recipe}


# ================================================================== workbook round trip
@safe
def load_workbook(name, data):
    blob, now = snapshot.read_blob(bytes(data))
    settings = blob.get("settings")
    global S
    remembered = S.project.remembered
    teachers = S.teachers
    S = Session(settings)
    S.project.remembered = remembered
    S.teachers = teachers or blob.get("teachers") or []
    ds, changes = snapshot.restore(blob, now, S.cfg, name)
    S.ds = ds
    S.recipe = blob.get("recipe") or {"items": []}
    S.restored = name
    _compute()
    d = summary_dict()
    d["changes_in_excel"] = changes
    return d


# ================================================================== double-entry check
@safe
def compare_entries(id_a, id_b):
    """Two people typed the same marks: list every cell where they differ."""
    p = S.project
    srcs = {s.id: s for s in p.sources}
    if id_a not in srcs or id_b not in srcs:
        raise ValueError("Pick two uploaded files.")

    def marks_of(sid):
        sub = Project(S.cfg)
        sub.sources = [srcs[sid]]
        sub.layouts = {k: v for k, v in p.layouts.items() if k[0] == sid}
        b = Builder(sub)
        ds = b.run()
        out = {}
        for m in ds.marks:
            st = ds.students[m.sid]
            ident = st.adm_no or st.board_roll or f"{st.cls}-{st.section}-{st.roll}" if (st.adm_no or st.board_roll or st.roll) else st.name.lower()
            out[(ident, m.exam, m.subject, m.component)] = (m.raw, st.name, m.ref)
        return out
    a, b = marks_of(id_a), marks_of(id_b)
    rows = []
    for k in sorted(set(a) | set(b)):
        va, vb = a.get(k), b.get(k)
        ta = va[0] if va else "(missing)"
        tb = vb[0] if vb else "(missing)"
        if ta.strip() != tb.strip():
            name = (va or vb)[1]
            rows.append({"student": name, "id": k[0], "exam": S.cfg.exam_name(k[1]), "subject": S.cfg.subject_name(k[2]),
                         "part": k[3], "a": ta, "b": tb,
                         "where_a": va[2].text() if va else "", "where_b": vb[2].text() if vb else ""})
    return {"a": srcs[id_a].name, "b": srcs[id_b].name, "checked": len(set(a) | set(b)), "differences": rows}


# ================================================================== on-device AI support
@safe
def intent_examples():
    from .nlu import _load_intents
    d = _load_intents()
    out = []
    for it in d["intents"]:
        for ex in it["examples"]:
            out.append({"intent": it["id"], "text": ex, "kind": "question"})
    for it in d["commands"]:
        for ex in it["examples"]:
            out.append({"intent": it["id"], "text": ex, "kind": "command"})
    return out


@safe
def header_labels():
    """Phrases for suggesting what an unknown column is (used by the on-device AI)."""
    cfg = S.cfg
    out = []
    for role, words in (("name", ["student name", "name of the student"]), ("father", ["father's name"]),
                        ("mother", ["mother's name"]), ("roll", ["roll number"]), ("adm_no", ["admission number"]),
                        ("board_roll", ["board roll number"]), ("section", ["section"]), ("class", ["class"]),
                        ("file_total", ["total marks"]), ("file_percent", ["percentage"]), ("file_grade", ["grade"]),
                        ("file_result", ["result pass or fail"]), ("ignore", ["rank", "remarks"])):
        for w in words:
            out.append({"label": w, "role": role})
    for code, s in cfg.subjects.items():
        out.append({"label": f"{s['name']} marks", "role": "mark", "subject": code, "component": "theory"})
        if s.get("practical"):
            out.append({"label": f"{s['name']} theory", "role": "mark", "subject": code, "component": "theory"})
            out.append({"label": f"{s['name']} practical", "role": "mark", "subject": code, "component": "practical"})
    return out
