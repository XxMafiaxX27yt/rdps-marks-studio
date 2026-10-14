"""Saves the clean data inside the workbook (hidden '_meta' sheet) so a teacher
can upload the workbook again later - with marks corrected in Excel - and keep
working: ask questions, add charts, rebuild. Marks typed in the Marks Matrix
win over the saved copy, and every such change is listed."""
import json
from decimal import Decimal

from .model import Student, Mark, Ref, Issue
from .build import Dataset
from .textutil import parse_number, clean, fmt_mark

VERSION = 1


def dump(sw):
    ds = sw.ds
    cells = []
    for (sid, exam), mrow in sw.mx_row_of.items():
        er = sw.res.exam_results[(sid, exam)]
        for row in er.rows:
            for comp in ("theory", "practical"):
                col = sw.col_of.get((row.subject, comp))
                if col is None:
                    continue
                if comp == "practical" and (row.is_total or not row.pr_max):
                    continue
                cells.append([mrow, col, sid, exam, row.subject, "total" if row.is_total else comp])
    return {
        "version": VERSION,
        "generated": sw.generated,
        "settings": sw.cfg.settings,
        "recipe": sw.recipe,
        "students": [s.to_dict() for s in ds.ordered_students()],
        "marks": [[m.sid, m.exam, m.subject, m.component, m.raw, m.kind, m.code,
                   None if m.value is None else str(m.value), m.ref.to_dict() if m.ref else None,
                   m.original, m.problem] for m in ds.marks],
        "cells": cells,
        "sources": [[a, b, len(rows), skipped] for a, b, rows, skipped in ds.raw_sheets],
    }


def is_ours(sheetnames):
    return "_meta" in sheetnames and "Marks Matrix" in sheetnames


def read_blob(data):
    """-> the saved JSON (dict) from an uploaded workbook, plus the current Marks Matrix values."""
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_vba=False)
    ws = wb["_meta"]
    if ws.cell(1, 1).value != "RDPS-MARKS-STUDIO":
        raise ValueError("This workbook was not made by RDPS Marks Studio.")
    parts = []
    r = 2
    while ws.cell(r, 1).value:
        parts.append(ws.cell(r, 1).value)
        r += 1
    blob = json.loads("".join(parts))
    mx = wb["Marks Matrix"]
    now = {}
    for mrow, col, sid, exam, subject, comp in blob["cells"]:
        v = mx.cell(mrow + 1, col + 1).value
        now[(sid, exam, subject, comp)] = v
    return blob, now


def restore(blob, now, cfg, file_name):
    """Rebuild the Dataset; marks changed in Excel are applied and reported."""
    ds = Dataset(cfg)
    for d in blob["students"]:
        st = Student(sid=d["sid"], name=d["name"], cls=d["cls"], section=d["section"], roll=d["roll"],
                     adm_no=d["adm_no"], board_roll=d["board_roll"], father=d["father"], mother=d["mother"],
                     gender=d.get("gender", ""), group=d.get("group", ""), group_given=d.get("group_given", ""),
                     group_note=d.get("group_note", ""), subjects=d["subjects"], choices_given=d.get("choices_given", False))
        st.refs = [Ref(r["file"], r["sheet"], r["row"], r["col"]) for r in d.get("refs", [])] or [Ref(file_name)]
        ds.students[st.sid] = st
    changes = 0
    for sid, exam, subject, comp, raw, kind, code, value, ref, original, problem in blob["marks"]:
        m = Mark(sid=sid, exam=exam, subject=subject, component=comp, raw=raw, kind=kind, code=code,
                 value=None if value is None else Decimal(value),
                 ref=Ref(ref["file"], ref["sheet"], ref["row"], ref["col"]) if ref else None,
                 original=original, problem=problem)
        key = (sid, exam, subject, comp)
        if key in now:
            cur = now[key]
            cur_txt = "" if cur is None else (fmt_mark(parse_number(cur)) if parse_number(cur) is not None else clean(cur))
            saved_txt = _saved_cell_text(m)
            if cur_txt != saved_txt:
                changes += 1
                _apply(m, cur_txt, cfg)
                st = ds.students[sid]
                ds.issue("info", "changed_in_excel",
                         f"{st.name}: {cfg.subject_name(subject)} {comp} ({cfg.exam_name(exam)}) was changed in Excel "
                         f"from “{saved_txt}” to “{cur_txt}”.", Ref(file_name, "Marks Matrix"), sid=sid,
                         subject=subject, exam=exam)
        ds.marks.append(m)
    for m in ds.marks:
        if m.kind == "invalid":
            st = ds.students[m.sid]
            ds.issue("error", "invalid_mark", f"{st.name}: {cfg.subject_name(m.subject)} {m.component} "
                     f"({cfg.exam_name(m.exam)}) - {m.problem}. Not used until corrected.", m.ref, sid=m.sid,
                     subject=m.subject, exam=m.exam)
    ds.raw_sheets = [(a, b, [], s) for a, b, n, s in blob.get("sources", [])]
    return ds, changes


def _saved_cell_text(m):
    if m.kind == "number":
        return fmt_mark(m.value)
    if m.kind == "leave":
        return m.code
    if m.kind == "invalid":
        return f"{m.raw} ⚠"
    return ""


def _apply(m, text, cfg):
    m.original = _saved_cell_text(m)
    t = text.replace(" ⚠", "").strip()
    m.raw = t
    m.problem = ""
    m.code = ""
    m.value = None
    if t == "":
        m.kind = "blank"
        return
    code = cfg.leave_code(t)
    if code:
        m.kind, m.code = "leave", code
        return
    v = parse_number(t)
    comps = cfg.components(m.subject, m.exam)
    mx = cfg.subject_max(m.subject, m.exam) if m.component == "total" else comps.get(m.component)
    if v is None:
        m.kind, m.problem = "invalid", f"“{t}” is not a number or a leave code"
    elif v < 0 or (mx is not None and v > mx):
        m.kind, m.problem, m.value = "invalid", f"{t} is outside 0-{fmt_mark(mx)}", v
    else:
        m.kind, m.value = "number", v
